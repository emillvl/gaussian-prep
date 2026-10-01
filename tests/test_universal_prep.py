import copy
import json
from types import SimpleNamespace

import pytest

from conftest import ScriptedProvider, audit_result, mathematical_model, question, review
from gaussian_prep import cli
from gaussian_prep.corpus import Corpus
from gaussian_prep.export import build_docx, build_pdf, render_latex
from gaussian_prep.models import MathModel, Plan, Question, Slot
from gaussian_prep.pipeline import Pipeline, PipelineError
from gaussian_prep.prep import prep_profile
from gaussian_prep.prompts import role_prompt
from gaussian_prep.provider import ModelProvider, ProviderError
from gaussian_prep.recovery import recover_saved_run
from gaussian_prep.run_files import RunFiles
from gaussian_prep.verification import compute


def verbal_corpus(tmp_path):
    path = tmp_path / "verbal.json"
    path.write_text(json.dumps([dict(id="v1", topic="Information and Ideas", skill="Central idea",
        difficulty="easy", format="MCQ", passage="Mira tried again after each failure. She finally solved the puzzle.",
        question="Which word best describes Mira?", choices=["Persistent", "Careless", "Indifferent", "Forgetful"],
        answer="A")]), encoding="utf-8")
    return Corpus(path, prep_profile("SAT Verbal"))


class VerbalProvider(ScriptedProvider):
    def ask(self, role, payload, schema):
        self.calls.append((role, copy.deepcopy(payload)))
        if role == "planner":
            return Plan(title="Verbal practice", interpretation="One source question", assumptions=[], requirements=[],
                total_questions=1, module_count=1, minutes_per_module=[2], slots=[Slot(id="v1", module=1, position=1,
                domain="Information and Ideas", subskill="Central idea", difficulty="Easy", format="MCQ", objective="Infer a trait from the passage.")])
        if role == "generator":
            source = payload["assigned_source"]
            return Question(slot_id="v1", stem=source["question"], format="MCQ", answer="D", solution="Repeated attempts show persistence.",
                choices=[dict(label=k, text=v) for k, v in zip("ABCD", ["Inattentive", "Uninterested", "Forgetful", "Determined"])], figure=None, table=None)
        if role == "verifier":
            assert "answer" not in payload["question"] and "solution" not in payload["question"]
            assert "assigned_source" not in payload and "reference_examples" not in payload
            return MathModel(mode="unsupported", interpretation="Infer character from actions in the passage.", variables=[], equations=[], constraints=[],
                target="", aggregate="each", options=[], unsupported_reason="Reading comprehension needs textual evidence.",
                independent_answer="D", evidence="She kept trying after failure, which supports determination and rules out indifference and inattention.")
        if role == "reviewer":
            return review()
        return audit_result(payload, schema, accepted=True, summary="Source-based item checked", issues=[])


def test_verbal_keeps_all_five_roles_source_text_and_independent_answer(tmp_path):
    corpus = verbal_corpus(tmp_path)
    original = corpus.path.read_bytes()
    provider = VerbalProvider()
    state = Pipeline(corpus, provider, tmp_path / "run", compute=compute).run("Generate 1 question")
    assert [role for role, _ in provider.calls] == ["planner", "generator", "verifier", "reviewer", "auditor"]
    assert all(payload["prep_profile"]["name"] == "SAT Verbal" for _, payload in provider.calls)
    item = state["items"]["v1"]
    assert item["question"]["stem"] == corpus.references[0].question
    assert item["question"]["answer"] == "D"
    assert item["verification"]["independent_answer"] == "D"
    assert item["verification"]["status"] == "unverified"  # Never presented as computational proof.
    assert item["similarity"]["exact_source_copy"]
    assert state["plan"]["slots"][0]["source_index"] == 0
    assert corpus.path.read_bytes() == original
    recovered = recover_saved_run(corpus, tmp_path / "run", compute=compute)
    assert recovered["status"] == "complete"
    with pytest.raises(PipelineError, match="original prep material"):
        recover_saved_run(Corpus(corpus.path), tmp_path / "run", compute=compute)


@pytest.mark.parametrize("defect", ["invented_stem", "wrong_key", "missing_evidence", "unchanged_choices"])
def test_verbal_rejects_invention_and_unverified_answers(tmp_path, defect):
    class Bad(VerbalProvider):
        def ask(self, role, payload, schema):
            result = super().ask(role, payload, schema)
            if role == "generator" and defect == "invented_stem":
                result.stem = "A new invented passage. What does the writer imply?"
            if role == "generator" and defect == "unchanged_choices":
                data = result.model_dump()
                data["choices"] = [dict(label=k, text=v) for k, v in zip("ABCD", payload["assigned_source"]["choices"])]
                data["answer"] = "A"
                result = Question.model_validate(data)
            if role == "verifier" and defect == "wrong_key":
                result.independent_answer = "A"
            if role == "verifier" and defect == "missing_evidence":
                result.evidence = ""
            return result
    provider = Bad()
    with pytest.raises(PipelineError, match="did not pass"):
        Pipeline(verbal_corpus(tmp_path), provider, tmp_path / "run", revisions=0, compute=compute).run("Generate 1 question")
    assert not any(role == "auditor" for role, _ in provider.calls)


def test_source_allocation_never_invents_extra_questions(tmp_path):
    corpus = verbal_corpus(tmp_path)
    plan = VerbalProvider().ask("planner", {}, Plan)
    plan.slots.append(plan.slots[0].model_copy(update={"id": "v2", "position": 2}))
    with pytest.raises(ValueError, match="Not enough distinct source"):
        corpus.assign_sources(plan)


@pytest.mark.parametrize("name,source_based", [("SAT Math", False), ("SAT Reading and Writing", True),
    ("AP Calculus BC", False), ("A-Levels Mathematics", False), ("GCSE Physics", False),
    ("ACT English", True), ("CSCA Mathematics", False), ("AP History", True)])
def test_subject_profiles_select_scope(name, source_based):
    profile = prep_profile(name)
    assert profile.source_based == source_based
    for role in ("planner", "generator", "verifier", "reviewer", "auditor"):
        prompt = role_prompt(role, profile.to_dict())
        if profile.name != "SAT Math":
            assert 'generator of NEW SAT Math' not in prompt
            assert "Do not create calculus questions" not in prompt
            assert "44-question" not in prompt
            assert profile.name in prompt


def test_sat_requires_explicit_section(monkeypatch, capsys):
    answers = iter(["SAT", "SAT Math and SAT Verbal", "SAT Verbal"])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    assert cli.choose_prep_material().name == "SAT Verbal"
    assert capsys.readouterr().out.count("Please explicitly choose SAT Math or SAT Verbal") == 2


@pytest.mark.parametrize("count", [2, 5, 8])
def test_computation_checks_every_displayed_choice(count):
    data = question().model_dump()
    data["choices"] = [dict(label=k, text=str(i)) for i, k in enumerate("ABCDEFGH"[:count], 6)]
    q = Question.model_validate(data)
    model = mathematical_model().model_dump()
    model["options"] = [dict(label=c.label, expression=c.text) for c in q.choices]
    assert compute(q, MathModel.model_validate(model))["status"] == "pass"


def test_general_taxonomy_and_calculus_are_not_sat_restricted(tmp_path):
    path = tmp_path / "calculus.json"
    path.write_text(json.dumps([dict(question="Evaluate the integral of the function.", topic="Integration", skill="Definite integrals", format="FRQ", difficulty="Hard")]))
    corpus = Corpus(path, prep_profile("AP Calculus BC"))
    assert corpus.taxonomy == {"Integration": ["Definite integrals"]}
    assert "calculus_outside_scope" not in corpus.references[0].quality_notes
    assert corpus.references[0].format == "FRQ"


def test_csv_jsonl_and_wrapped_datasets_preserve_passages(tmp_path):
    row = dict(question="Which answer fits the blank ____?", passage="Some source context.", topic="Grammar", choices=["A", "B"], format="MCQ")
    for filename, text in [("set.json", json.dumps({"questions": [row]})), ("set.jsonl", json.dumps(row) + "\n"),
        ("set.csv", 'question,passage,topic,format,A,B\nWhich answer fits the blank ____?,Some source context.,Grammar,MCQ,A,B\n')]:
        path = tmp_path / filename
        path.write_text(text, encoding="utf-8")
        corpus = Corpus(path, prep_profile("ACT English"))
        assert corpus.references[0].question == row["passage"] + "\n\n" + row["question"]
        assert len(corpus.references[0].choices) == 2


def test_discovery_and_missing_dataset_open_chooser(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    for key in ("GAUSSIAN_DATASET", "SATPREP_DATASET"):
        monkeypatch.delenv(key, raising=False)
    selected = tmp_path / "picked.json"
    picks = []
    monkeypatch.setattr(cli, "pick_dataset", lambda: picks.append(True) or selected)
    assert cli.dataset_path(None) == selected
    assert "Please provide us with a dataset." in capsys.readouterr().out
    selected.write_text(json.dumps([{"question": "An existing source question?"}]))
    (tmp_path / "state.json").write_text('{"items": {}}')
    assert cli.dataset_path(None) == selected
    assert len(picks) == 1
    (tmp_path / "another.json").write_bytes(selected.read_bytes())
    assert cli.dataset_path(None) == selected
    assert len(picks) == 2


def test_picker_cancellation_closes_window(tmp_path, monkeypatch):
    import sys
    calls = []
    root = SimpleNamespace(withdraw=lambda: calls.append("hide"), attributes=lambda *a: None, destroy=lambda: calls.append("close"))
    monkeypatch.setitem(sys.modules, "tkinter", SimpleNamespace(Tk=lambda: root, filedialog=SimpleNamespace(askopenfilename=lambda **k: "")))
    with pytest.raises(ValueError, match="selection cancelled"):
        cli.pick_dataset()
    assert calls == ["hide", "close"]


def run_cli(corpus, provider, run_dir, monkeypatch, resume=False):
    monkeypatch.setattr(cli, "ModelProvider", lambda *a, **kw: provider)
    monkeypatch.setattr(cli, "resolve_credentials", lambda *a: ("opencode", provider.model, "test-key"))
    provider.name = "opencode"
    command = ["resume", str(run_dir)] if resume else ["generate", f"Generate {provider.count} questions", "--dataset", str(corpus.path), "--prep-material", corpus.profile.name, "--out", str(run_dir)]
    return cli.main(command + ["--format", "none"])


@pytest.mark.parametrize("failures", [1, 2, 3])
def test_run_automatically_retries_twice_without_recreating_accepted_items(corpus, tmp_path, monkeypatch, capsys, failures):
    class Failing(ScriptedProvider):
        attempts = 0
        def ask(self, role, payload, schema):
            if role == "generator" and payload["slot"]["id"] == "q2":
                self.attempts += 1
                if self.attempts <= failures:
                    raise ProviderError("Transient provider failure")
            return super().ask(role, payload, schema)
    provider = Failing(count=2)
    run = tmp_path / "run"
    assert run_cli(corpus, provider, run, monkeypatch) == (1 if failures == 3 else 0)
    assert provider.attempts == min(3, failures + 1)
    assert provider.generated["q1"] == 1
    output = capsys.readouterr().out
    assert output.count("Automatic retry ") == min(2, failures)
    state = json.loads(RunFiles(run).checkpoint.read_text())
    assert len(state["automatic_retries"]) == min(2, failures)
    if failures == 3:
        assert "Continue when ready:" in output
        assert state["status"] == "needs_attention"
        assert run_cli(corpus, provider, run, monkeypatch, resume=True) == 0
        assert provider.generated["q1"] == 1


def test_final_audit_retries_only_the_four_pending_questions(corpus, tmp_path, monkeypatch):
    class FinalAudit(ScriptedProvider):
        followups = 0
        def ask(self, role, payload, schema):
            if role == "auditor":
                self.calls.append((role, copy.deepcopy(payload)))
                if payload["audit_mode"] == "initial":
                    return audit_result(payload, schema, accepted=False, summary="Four questions need revision", issues=[dict(
                        dimension="variety", reason="Revise four constructions", slot_ids=["q2", "q3", "q4", "q5"], instruction="Use distinct constructions")])
                self.followups += 1
                if self.followups < 3:
                    raise ProviderError("Final audit unavailable")
                return audit_result(payload, schema, accepted=True, summary="All four checked", issues=[])
            if role == "generator" and self.generated.get(payload["slot"]["id"]):
                sid = payload["slot"]["id"]
                self.generated[sid] += 1
                return question(int(sid[1:]), 101 + 2 * int(sid[1:]))
            return super().ask(role, payload, schema)
    provider = FinalAudit(count=5)
    run = tmp_path / "run"
    assert run_cli(corpus, provider, run, monkeypatch) == 0
    audits = [p for r, p in provider.calls if r == "auditor"]
    assert len(audits) == 4 and provider.followups == 3
    assert all(p["scope_slot_ids"] == ["q2", "q3", "q4", "q5"] for p in audits[1:])
    assert provider.generated == {"q1": 1, "q2": 2, "q3": 2, "q4": 2, "q5": 2}


def test_non_sat_exports_use_subject_and_written_answers(tmp_path):
    corpus = verbal_corpus(tmp_path)
    state = Pipeline(corpus, VerbalProvider(), tmp_path / "run", compute=compute).run("Generate 1 question")
    state["prep_profile"] = prep_profile("AP History").to_dict()
    state["plan"]["slots"][0]["format"] = "FRQ"
    state["plan"]["structure"]["sections"][0]["format_sequence"] = ["FRQ"]
    state["items"]["v1"]["question"].update(format="FRQ", choices=[], answer="The evidence supports persistence.")
    build_docx(state, tmp_path / "questions.docx")
    build_pdf(state, tmp_path / "questions.pdf")
    build_docx(state, tmp_path / "key.docx", answers=True)
    from docx import Document
    text = "\n".join(p.text for p in Document(tmp_path / "questions.docx").paragraphs)
    assert "AP History" in text and "Calculator permitted" not in text and "SAT-style" not in text
    assert "The evidence supports persistence." in "\n".join(p.text for p in Document(tmp_path / "key.docx").paragraphs)
    tex = render_latex(state)
    assert "AP History" in tex and "Calculator permitted" not in tex
    assert r"\newline{}\newline{}Which word" in tex


def test_biology_written_response_uses_the_same_five_roles(tmp_path):
    path = tmp_path / "biology.json"
    path.write_text(json.dumps([dict(question="Describe the role of the cell membrane.", topic="Cells", skill="Membranes", format="FRQ", difficulty="Easy")]))
    corpus = Corpus(path, prep_profile("AP Biology"))
    class Biology(ScriptedProvider):
        def ask(self, role, payload, schema):
            self.calls.append((role, copy.deepcopy(payload)))
            if role == "planner":
                return Plan(title="Biology practice", interpretation="Written response", assumptions=[], requirements=[],
                    total_questions=1, module_count=1, minutes_per_module=[5], slots=[Slot(id="b1", module=1, position=1,
                    domain="Cells", subskill="Membranes", difficulty="Easy", format="FRQ", objective="Explain selective permeability.")])
            if role == "generator":
                return Question(slot_id="b1", stem="Explain why a cell membrane allows some substances to cross more readily than others.", format="FRQ",
                    choices=[], answer="The lipid bilayer and membrane proteins provide selective permeability.", solution="Discuss hydrophobic core and transport proteins.", figure=None, table=None)
            if role == "verifier":
                assert "answer" not in payload["question"]
                return MathModel(mode="unsupported", interpretation="Biological explanation", variables=[], equations=[], constraints=[], target="", aggregate="each",
                    options=[], unsupported_reason="Requires a written biological explanation.", independent_answer="The hydrophobic core restricts charged substances; specific proteins mediate transport.",
                    evidence="The phospholipid bilayer has a hydrophobic interior. Channels and carriers are selective.")
            if role == "reviewer":
                return review()
            return audit_result(payload, schema, accepted=True, summary="Written response checked", issues=[])
    provider = Biology()
    state = Pipeline(corpus, provider, tmp_path / "run", compute=compute).run("Generate 1 question")
    assert state["status"] == "complete"
    assert [role for role, _ in provider.calls] == ["planner", "generator", "verifier", "reviewer", "auditor"]


def test_profile_changes_actual_provider_system_instructions(tmp_path):
    provider = ModelProvider("test-model", "test-key", tmp_path)
    _, _, body, _ = provider.request_parts("generator", {"prep_profile": prep_profile("SAT Verbal").to_dict()}, Question, [])
    prompt = body["messages"][0]["content"]
    assert "SOURCE-BASED MODE" in prompt and "SAT Verbal" in prompt
    assert "Do not copy a question" not in prompt


def test_request_examples_are_universal_placeholders():
    examples = dict(cli.example_requests())
    assert list(examples) == ["Mock test", "Practice test"]
    mock, practice = examples["Mock test"], examples["Practice test"]
    assert "[TOPIC]" in mock and "[SUBTOPIC" not in mock
    assert "[TOPIC]" in practice and "[SUBTOPIC1]" in practice and "[SUBTOPIC2]" in practice
    assert all("SAT" not in text and "AP " not in text for text in examples.values())


def test_resume_relocates_original_dataset_with_picker(corpus, tmp_path, monkeypatch):
    run = tmp_path / "run"
    state = Pipeline(corpus, ScriptedProvider(), run, compute=compute).run("Generate 1 question")
    relocated = tmp_path / "relocated.json"
    relocated.write_bytes(corpus.path.read_bytes())
    corpus.path.unlink()
    monkeypatch.setattr(cli, "pick_dataset", lambda: relocated)
    monkeypatch.setattr(cli, "resolve_credentials", lambda *a: pytest.fail("Completed resume must remain offline"))
    assert cli.main(["resume", str(run), "--format", "none"]) == 0
    saved = json.loads(RunFiles(run).checkpoint.read_text())
    assert saved["dataset"] == str(relocated)
    assert saved["items"] == state["items"]


@pytest.mark.parametrize("interruption", ["conflict", "keyboard"])
def test_user_action_cases_are_not_automatically_retried(corpus, tmp_path, monkeypatch, capsys, interruption):
    class Stop(ScriptedProvider):
        attempts = 0
        def ask(self, role, payload, schema):
            self.attempts += 1
            if interruption == "keyboard":
                raise KeyboardInterrupt()
            return Plan(title="Conflicting request", interpretation="Needs clarification", assumptions=[], requirements=[], conflicts=["Please choose one topic."],
                total_questions=0, module_count=0, slots=[], minutes_per_module=[])
    provider = Stop()
    assert run_cli(corpus, provider, tmp_path / "run", monkeypatch) == (130 if interruption == "keyboard" else 1)
    assert provider.attempts == 1
    assert "Automatic retry" not in capsys.readouterr().out
