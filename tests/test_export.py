import zipfile

import pytest

from gaussian_prep.export import _escape_latex, _segments, build_docx, build_latex_pdf, build_pdf, export, figure_png, find_latex, math_png, render_latex
from gaussian_prep.models import Figure, Plan, Question, Slot


def state():
    slot = Slot(id="q1", module=1, position=1, domain="Algebra", subskill="Linear equations in one variable", difficulty="Easy", format="MCQ", objective="Solve a linear equation.")
    question = Question(
        slot_id="q1",
        stem=r"Find the value of \(x^2 + \frac{1}{2}\) when \(x = 3\).",
        format="MCQ",
        choices=[{"label": label, "text": text} for label, text in zip("ABCD", [r"\(9\frac{1}{2}\)", "10", "8", r"\(9\frac{3}{2}\)"])],
        answer="A",
        solution="Substitute x = 3 into the expression.",
        table={"headers": ["x", "f(x)"], "rows": [["1", "2"], ["2", "5"]]},
        figure=Figure(description="Right triangle", show_axes=False, points=[{"label": "A", "x": 0, "y": 0}, {"label": "B", "x": 3, "y": 0}, {"label": "C", "x": 0, "y": 4}], segments=[{"start": "A", "end": "B"}, {"start": "B", "end": "C"}, {"start": "C", "end": "A"}], circles=[]),
    )
    plan = Plan(title="Sample Test", interpretation="Fixture.", assumptions=[], requirements=[], total_questions=1, module_count=1, minutes_per_module=[10], slots=[slot])
    return {"status": "complete", "plan": plan.model_dump(), "items": {"q1": {"question": question.model_dump()}}}


def test_segments_split_text_and_math():
    parts = _segments(r"Solve \(x+1=2\) now")
    assert parts == [("text", "Solve "), ("math", "x+1=2"), ("text", " now")]


def test_math_image_and_bad_input():
    assert math_png(r"x^2+\frac{1}{2}").startswith(b"\x89PNG")
    with pytest.raises(ValueError, match="typeset unchanged"):
        math_png(r"\begin{nonsense}")


def test_figure_image_is_png():
    figure = Figure(description="Triangle", show_axes=False, points=[{"label": "A", "x": 0, "y": 0}], segments=[], circles=[])
    assert figure_png(figure).startswith(b"\x89PNG")


def test_builds_pdf_and_docx(tmp_path):
    pdf = tmp_path / "questions.pdf"
    docx = tmp_path / "questions.docx"
    build_pdf(state(), pdf)
    build_docx(state(), docx)
    assert pdf.read_bytes().startswith(b"%PDF") and pdf.stat().st_size > 1000
    assert zipfile.is_zipfile(docx) and docx.stat().st_size > 1000


def test_answer_key_has_no_question_text(tmp_path):
    key = tmp_path / "answer-key.pdf"
    build_pdf(state(), key, answers=True)
    assert key.read_bytes().startswith(b"%PDF")


def _install_logo(monkeypatch, tmp_path):
    from io import BytesIO
    from PIL import Image
    import gaussian_prep.export as module
    logo = tmp_path / "logo.png"
    buffer = BytesIO()
    Image.new("RGB", (40, 40), "white").save(buffer, "PNG")
    logo.write_bytes(buffer.getvalue())
    monkeypatch.setattr(module, "LOGO_PATH", logo)
    return module


def test_logo_page_is_appended_when_artwork_is_present(tmp_path, monkeypatch):
    module = _install_logo(monkeypatch, tmp_path)

    source = render_latex(state())
    assert r"\includegraphics" in source
    assert source.rstrip().endswith(r"\end{document}")
    assert source.index(r"\includegraphics") > source.index(r"\begin{enumerate}")

    captured = []
    monkeypatch.setattr(module.SimpleDocTemplate, "build", lambda self, story: captured.extend(story))
    build_pdf(state(), tmp_path / "questions.pdf")
    assert any(isinstance(flowable, module.RLImage) for flowable in captured)

    docx_path = tmp_path / "questions.docx"
    build_docx(state(), docx_path)
    from docx import Document
    from docx.oxml.ns import qn
    document = Document(docx_path)
    vAlign = document.sections[-1]._sectPr.find(qn("w:vAlign"))
    assert vAlign is not None and vAlign.get(qn("w:val")) == "center"
    assert document.element.body.xpath(".//w:drawing")


def test_no_logo_page_without_artwork(tmp_path, monkeypatch):
    import gaussian_prep.export as module
    monkeypatch.setattr(module, "LOGO_PATH", tmp_path / "missing.png")
    assert r"\includegraphics" not in render_latex(state())


def test_export_returns_written_paths(tmp_path):
    written = export(state(), tmp_path, "both")
    assert {path.name for path in written} == {"questions.pdf", "answer-key.pdf", "questions.docx", "answer-key.docx"}
    assert all(path.is_file() for path in written)


def test_export_refuses_incomplete_state(tmp_path):
    with pytest.raises(ValueError):
        export({"status": "needs_attention"}, tmp_path, "pdf")


def test_latex_source_has_math_table_and_tikz():
    source = render_latex(state())
    assert r"\(x^2 + \frac{1}{2}\)" in source
    assert "\\begin{tikzpicture}" in source
    assert "anchor=" in source
    assert "\\begin{tabular}" in source
    assert "\\begin{enumerate}" in source
    assert source.rstrip().endswith("\\end{document}")


def test_latex_answer_key_hides_question_text():
    source = render_latex(state(), answers=True)
    assert "Find the value" not in source
    assert "\\textbf{1.}" in source


def test_latex_escapes_reserved_characters():
    assert _escape_latex("100% & $5_0 #1") == r"100\% \& \$5\_0 \#1"
    assert _escape_latex("90\u00b0") == r"90\ensuremath{^\circ}"
    assert _escape_latex("Caf\u00e9") == "Caf\u00e9"


def test_latex_maps_unicode_and_stays_ascii():
    st = state()
    st["items"]["q1"]["question"]["stem"] = "Find \u03b8 if \u2220A = 90\u00b0 and x \u2264 5 (\u221a2 \u2248 1.4)."
    source = render_latex(st)
    assert source.isascii()
    assert r"\ensuremath{\theta}" in source
    assert r"\ensuremath{\le}" in source
    assert r"\ensuremath{^\circ}" in source


def test_latin_safe_transliterates_for_reportlab():
    from gaussian_prep.export import _latin_safe
    assert _latin_safe("x \u2264 5 \u2212 2\u03c0") == "x <= 5 - 2pi"
    assert _latin_safe("caf\u00e9 \u00b0") == "caf\u00e9 \u00b0"


def test_answer_latex_renders_engine_notation():
    from gaussian_prep.export import _answer_latex
    assert _answer_latex("B") == "B"
    assert _answer_latex("3/4") == r"\(\frac{3}{4}\)"
    assert _answer_latex("sqrt(2/3)") == r"\(\frac{\sqrt{6}}{3}\)"
    assert _answer_latex("2**3") == r"\(8\)"
    assert _answer_latex("pi") == r"\(\pi\)"
    assert _answer_latex("2**10") == r"\(1024\)"
    assert _answer_latex("sqrt(1+sqrt(2))") == r"\(\sqrt{1 + \sqrt{2}}\)"
    assert _answer_latex("2*pi") == r"\(2 \pi\)"


def test_display_math_stays_display_math():
    from gaussian_prep.export import _latex_text
    assert _latex_text(r"Evaluate \[\frac{1}{2}\].") == r"Evaluate \[\frac{1}{2}\]."


def test_docx_table_math_is_typeset(tmp_path):
    st = state()
    st["items"]["q1"]["question"]["table"] = {"headers": [r"\(x^2\)"], "rows": [[r"\(\frac{1}{2}\)"]]}
    path = tmp_path / "table.docx"
    build_docx(st, path)
    from docx import Document
    table = Document(path).tables[0]
    assert "\\(" not in table.cell(0, 0).text
    assert len(table._element.xpath(".//w:drawing")) == 2


@pytest.mark.parametrize("failure", [RuntimeError("compile failed"), OSError("cannot start compiler")])
def test_export_falls_back_on_compiler_failure(tmp_path, monkeypatch, failure):
    import gaussian_prep.export as module
    monkeypatch.setattr(module, "find_latex", lambda: "pdflatex")
    def fail(*args, **kwargs):
        raise failure
    monkeypatch.setattr(module, "build_latex_pdf", fail)
    written = export(state(), tmp_path)
    assert all(path.read_bytes().startswith(b"%PDF") for path in written)
    assert (tmp_path / "support" / "sources" / "questions.tex").is_file()


def test_export_falls_back_on_compiler_timeout(tmp_path, monkeypatch):
    import subprocess
    test_export_falls_back_on_compiler_failure(tmp_path, monkeypatch, subprocess.TimeoutExpired("pdflatex", 300))


def test_export_without_compiler_keeps_latex_source(tmp_path, monkeypatch):
    monkeypatch.setattr("gaussian_prep.export.find_latex", lambda: None)
    export(state(), tmp_path)
    assert (tmp_path / "support" / "sources" / "questions.tex").is_file()
    assert (tmp_path / "support" / "sources" / "answer-key.tex").is_file()


def test_find_latex_ignores_inaccessible_optional_paths(monkeypatch):
    from pathlib import Path
    monkeypatch.setattr("gaussian_prep.export.shutil.which", lambda name: None)
    def inaccessible(path):
        raise PermissionError("Cannot inspect optional installation")
    monkeypatch.setattr(Path, "is_file", inaccessible)
    assert find_latex() is None


def test_tectonic_uses_its_own_command_line(tmp_path, monkeypatch):
    import gaussian_prep.export as module
    from types import SimpleNamespace
    monkeypatch.setattr(module, "find_latex", lambda: "tectonic")
    commands = []
    def run(command, **kwargs):
        commands.append(command)
        (tmp_path / "questions.pdf").write_bytes(b"%PDF test")
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(module.subprocess, "run", run)
    build_latex_pdf(state(), tmp_path / "questions.pdf")
    assert commands[0][1:3] == ["--outdir", str(tmp_path)]


def test_latex_does_not_double_escape_math_percent():
    from gaussian_prep.export import _latex_text
    assert _latex_text(r"\(54\%\)") == r"\(54\%\)"
    with pytest.raises(ValueError, match="Escape literal"):
        _latex_text(r"\(50%\)")


@pytest.mark.skipif(find_latex() is None, reason="No LaTeX engine installed")
def test_latex_compiles_when_engine_available(tmp_path):
    st = state()
    st["items"]["q1"]["question"]["stem"] = "Find \u03b8 when \u2220A = 90\u00b0 and \\(x \\leq 5\\) (\u2248 54%)."
    build_latex_pdf(st, tmp_path / "questions.pdf")
    assert (tmp_path / "questions.pdf").read_bytes().startswith(b"%PDF")
    assert (tmp_path / "questions.tex").is_file()
    assert not (tmp_path / "questions.log").exists()


def test_fallback_allocates_enough_line_height_for_tall_math(tmp_path, monkeypatch):
    from reportlab.platypus import KeepTogether, Paragraph
    import gaussian_prep.export as module
    st = state()
    expression = r"\frac{\sqrt{1+\sqrt{2}}}{\frac{3}{4}}"
    st['items']['q1']['question']['choices'][0]['text'] = r'\(' + expression + r'\)'
    captured = []
    monkeypatch.setattr(module.SimpleDocTemplate, 'build', lambda self, story: captured.extend(story))
    build_pdf(st, tmp_path / 'tall-math.pdf')
    choices = [p for block in captured if isinstance(block, KeepTogether) for p in block._content if isinstance(p, Paragraph) and p.style.name == 'choice']
    _, allocated_height = choices[0].wrap(module._PDF_CONTENT_WIDTH, 1000)
    _, pixels = module._png_size(math_png(expression, fontsize=11))
    assert allocated_height >= pixels * 72 / 200 - 0.1


def test_fallback_wide_math_and_table_fit_the_page(tmp_path):
    st = state()
    wide = r'\(' + '+'.join(f'x_{{{i}}}' for i in range(50)) + r'\)'
    st['items']['q1']['question']['choices'][0]['text'] = wide
    st['items']['q1']['question']['table'] = {'headers':['Expression','Description'], 'rows':[[wide,'A long expression that must fit within its own table cell.']]}
    path = tmp_path / 'wide-math.pdf'
    build_pdf(st, path)
    assert path.read_bytes().startswith(b'%PDF')
