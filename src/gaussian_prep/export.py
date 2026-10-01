"""Render an accepted test to PDF and DOCX question sheets (answers-only key)."""
from __future__ import annotations

from html import escape
from io import BytesIO
import os
from pathlib import Path
import re
import shutil
import subprocess

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Image as RLImage
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import AssessmentStructure, Figure, Plan, Question
from .expressions import ExpressionReader, MathUnsupported
from .render import MATH, question_html, text_html
from .typesetting import RENDER_LOCK, math_png
from .run_files import RunFiles
from .prep import saved_profile
from .structure import assessment_label, check_structure_request, section_for, section_name, structure_errors

# Letter paper minus document margins and ReportLab's two frame paddings.
_PDF_CONTENT_WIDTH = 6.9 * inch - 12

LOGO_PATH = Path(__file__).with_name("logo.png")
LOGO_NAME = "logo.png"
LOGO_DISPLAY_INCHES = 4.0


def _logo_size() -> tuple[float, float] | None:
    """Return the logo's (width, height) in points, or None when no artwork is installed."""
    if not LOGO_PATH.is_file():
        return None
    width, height = _png_size(LOGO_PATH.read_bytes())
    display_width = LOGO_DISPLAY_INCHES * inch
    return display_width, display_width * height / width


class LatexUnavailable(RuntimeError):
    pass


# reportlab's base-14 fonts only carry Latin-1; transliterate common math symbols for that fallback path.
_PLAIN = {"\u2212": "-", "\u2264": "<=", "\u2265": ">=", "\u2260": "!=", "\u2248": "~=", "\u221e": "inf", "\u221a": "sqrt", "\u03c0": "pi", "\u03b8": "theta", "\u03b1": "alpha", "\u03b2": "beta", "\u03b3": "gamma", "\u0394": "Delta", "\u2220": "angle", "\u22a5": "perp", "\u2225": "||", "\u00b2": "^2", "\u00b3": "^3", "\u00b9": "^1", "\u00bd": "1/2", "\u00bc": "1/4", "\u00be": "3/4", "\u2192": "->", "\u21d2": "=>", "\u2208": " in ", "\u00b7": "*", "\u2009": " ", "\u202f": " ", "\u00a0": " ", "\u2013": "-", "\u2014": "--", "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2026": "..."}


def _latin_safe(text: str) -> str:
    return "".join(_PLAIN.get(character, character) for character in text if ord(character) < 256 or character in _PLAIN)


def figure_png(figure: Figure) -> bytes:
    points = {p.label: p for p in figure.points}
    xs = [p.x for p in figure.points] + [v for c in figure.circles for v in (c.x - c.radius, c.x + c.radius)]
    ys = [p.y for p in figure.points] + [v for c in figure.circles for v in (c.y - c.radius, c.y + c.radius)]
    if figure.show_axes:
        xs.append(0)
        ys.append(0)
    xmin, xmax, ymin, ymax = min(xs), max(xs), min(ys), max(ys)
    width, height = 440, 320
    scale = min((width - 40) / max(xmax - xmin, 1), (height - 40) / max(ymax - ymin, 1))
    xmid, ymid = (xmin + xmax) / 2, (ymin + ymax) / 2

    def sx(value):
        return width / 2 + (value - xmid) * scale

    def sy(value):
        return height / 2 + (value - ymid) * scale

    # Share Matplotlib state with formula rendering; labels remain literal on a fixed canvas.
    with RENDER_LOCK, matplotlib.rc_context({"text.usetex": False, "text.parse_math": False, "savefig.bbox": None, "savefig.pad_inches": 0}):
        fig, ax = plt.subplots(figsize=(width / 100, height / 100), dpi=200)
        try:
            ax.set_xlim(0, width)
            ax.set_ylim(0, height)
            ax.set_aspect("equal")
            ax.axis("off")
            if figure.show_axes:
                ax.plot([20, width - 20], [sy(0), sy(0)], color="#aab1b7", linewidth=1)
                ax.plot([sx(0), sx(0)], [20, height - 20], color="#aab1b7", linewidth=1)
            for segment in figure.segments:
                start, end = points[segment.start], points[segment.end]
                ax.plot([sx(start.x), sx(end.x)], [sy(start.y), sy(end.y)], color="#20252a", linewidth=1.8)
            for circle in figure.circles:
                ax.add_patch(plt.Circle((sx(circle.x), sy(circle.y)), circle.radius * scale, fill=False, edgecolor="#20252a", linewidth=1.8))
            centroid_x = sum(sx(p.x) for p in figure.points) / len(figure.points) if figure.points else width / 2
            centroid_y = sum(sy(p.y) for p in figure.points) / len(figure.points) if figure.points else height / 2
            for point in figure.points:
                x, y = sx(point.x), sy(point.y)
                dx, dy = x - centroid_x, y - centroid_y
                norm = (dx * dx + dy * dy) ** 0.5 or 1.0
                ax.plot(x, y, "o", color="#20252a", markersize=3)
                # Place each letter outwards from the figure so close points do not overlap.
                ax.annotate(point.label, (x, y), textcoords="offset points", xytext=(11 * dx / norm, 11 * dy / norm), ha="center", va="center", fontsize=9, family="serif", parse_math=False, usetex=False, bbox=dict(boxstyle="round,pad=0.14", fc="white", ec="none", alpha=0.8))
            ax.margins(0.14)
            buffer = BytesIO()
            fig.savefig(buffer, format="png", dpi=200, bbox_inches=None, transparent=True)
            return buffer.getvalue()
        finally:
            plt.close(fig)


def _png_size(png: bytes) -> tuple[int, int]:
    with Image.open(BytesIO(png)) as image:
        return image.size


def _segments(text: str):
    text_html(text)  # Validate without rewriting the author's source.
    parts, end = [], 0
    for match in MATH.finditer(text):
        if match.start() > end:
            parts.append(("text", text[end:match.start()]))
        parts.append(("math", match.group(1) if match.group(1) is not None else match.group(2)))
        end = match.end()
    if end < len(text):
        parts.append(("text", text[end:]))
    return parts


def _ordered(state: dict):
    plan = Plan.model_validate(state["plan"])
    dataset_structure = state.get("dataset_exam_structure")
    check_structure_request(plan, state.get("request", ""), saved_profile(state).name,
        AssessmentStructure.model_validate(dataset_structure) if dataset_structure else None)
    errors = structure_errors(plan, state["items"])
    if any(issue.get("dimension") == "structure" for verdict in (state.get("audit"), state.get("audit_warnings")) if verdict for issue in verdict.get("issues", [])):
        errors.append("The auditor's exam-structure conflict is unresolved")
    if errors:
        raise ValueError("Cannot export a test that fails structure checks: " + "; ".join(errors))
    result = []
    for module in range(1, plan.module_count + 1):
        slots = sorted((s for s in plan.slots if s.module == module), key=lambda s: s.position)
        result.append((module, slots))
    return plan, result


def _question(state: dict, slot_id: str) -> Question:
    question = Question.model_validate(state["items"][slot_id]["question"])
    question_html(question, 1)
    return question


def build_pdf(state: dict, path: Path, answers: bool = False, image_dir: Path | None = None) -> None:
    plan, modules = _ordered(state)
    styles = {
        "title": ParagraphStyle("title", fontName="Times-Bold", fontSize=20, leading=24, spaceAfter=6),
        "subtitle": ParagraphStyle("subtitle", fontName="Times-Roman", fontSize=10, textColor=colors.HexColor("#55636d"), spaceAfter=14),
        "module": ParagraphStyle("module", fontName="Times-Bold", fontSize=14, leading=17, spaceBefore=14, spaceAfter=4),
        "meta": ParagraphStyle("meta", fontName="Times-Roman", fontSize=9, textColor=colors.HexColor("#55636d"), spaceAfter=8),
        "stem": ParagraphStyle("stem", fontName="Times-Roman", fontSize=11.5, leading=17, autoLeading="max", spaceAfter=6),
        "choice": ParagraphStyle("choice", fontName="Times-Roman", fontSize=11, leading=15, autoLeading="max", leftIndent=18, spaceAfter=2),
        "answer": ParagraphStyle("answer", fontName="Times-Roman", fontSize=12, leading=18, autoLeading="max", spaceAfter=4),
    }
    title = _latin_safe(plan.title + (" \u2014 Answer Key" if answers else ""))
    story = [Paragraph(escape(title), styles["title"]), Paragraph(escape(_latin_safe(f"{saved_profile(state).name} · {assessment_label(plan)} · {plan.total_questions} questions · {plan.module_count} section(s)")), styles["subtitle"])]
    workdir = image_dir if image_dir is not None else path.parent / f".{path.stem}-images"
    workdir.mkdir(parents=True, exist_ok=True)
    counter = 0
    for module, slots in modules:
        if not answers:
            story.append(Paragraph(escape(_latin_safe(section_name(plan, module))), styles["module"]))
            story.append(Paragraph(f"{len(slots)} questions \u00b7 {plan.minutes_per_module[module-1]} minutes", styles["meta"]))
            section = section_for(plan, module)
            for instruction in section.instructions if section else []:
                story.append(Paragraph(_pdf_rich(instruction, workdir, 12), styles["meta"]))
        for slot in slots:
            counter += 1
            if answers:
                question = _question(state, slot.id)
                story.append(Paragraph(f"{counter}. {_pdf_rich(_answer_text(question.answer), workdir, 15)}", styles["answer"]))
                continue
            story.append(_pdf_question(state, slot.id, counter, styles, workdir))
        if module != modules[-1][0]:
            story.append(PageBreak())
    story += _logo_page()
    document = SimpleDocTemplate(str(path), pagesize=letter, topMargin=0.75 * inch, bottomMargin=0.75 * inch, leftMargin=0.8 * inch, rightMargin=0.8 * inch, title=title)
    document.build(story)


def _logo_page() -> list:
    """A blank page carrying only the centered brand logo, or nothing when no artwork is installed."""
    size = _logo_size()
    if size is None:
        return []
    display_width, display_height = size
    content_height = letter[1] - 1.5 * inch - 12
    return [PageBreak(), Spacer(1, max(0.0, (content_height - display_height) / 2)), RLImage(str(LOGO_PATH), width=display_width, height=display_height, hAlign="CENTER")]


def _inline_image(expression: str, workdir: Path, height: float, max_width: float = _PDF_CONTENT_WIDTH) -> str:
    png = math_png(expression, fontsize=height - 2)
    width, tall = _png_size(png)
    target = workdir / f"math-{abs(hash((expression, round(height, 2))))}.png"
    if not target.is_file():
        target.write_bytes(png)
    # Preserve the typesetter's dimensions: forcing every expression to the same height
    # enlarges single variables and crushes fractions and nested radicals.
    point_width, point_height = width * 72 / 200, tall * 72 / 200
    if point_width > max_width:
        point_height *= max_width / point_width
        point_width = max_width
    return f'<img src="{target.as_posix()}" width="{point_width:.1f}" height="{point_height:.1f}" valign="middle"/>'


def _pdf_rich(text: str, workdir: Path, height: float, max_width: float = _PDF_CONTENT_WIDTH) -> str:
    parts = []
    for kind, value in _segments(text):
        if kind == "text":
            parts.append(escape(_latin_safe(value)).replace("\n", "<br/>"))
        else:
            parts.append(_inline_image(value, workdir, height, max_width))
    return "".join(parts)


def _pdf_question(state: dict, slot_id: str, number: int, styles: dict, workdir: Path):
    question = _question(state, slot_id)
    flowables = [Paragraph(f"<b>{number}.</b>&nbsp;&nbsp;{_pdf_rich(question.stem, workdir, 14.5)}", styles["stem"])]
    if question.table:
        width = _PDF_CONTENT_WIDTH / len(question.table.headers)
        cell_style = ParagraphStyle("table-cell", parent=styles["choice"], leftIndent=0)
        data = [[Paragraph(_pdf_rich(cell, workdir, 13, width - 12), cell_style) for cell in question.table.headers]]
        data += [[Paragraph(_pdf_rich(cell, workdir, 13, width - 12), cell_style) for cell in row] for row in question.table.rows]
        table = Table(data, colWidths=[width] * len(question.table.headers), hAlign="CENTER", repeatRows=1)
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8294a0")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        flowables.append(Spacer(1, 6))
        flowables.append(table)
    if question.figure:
        png, path_ = figure_png(question.figure), workdir / f"figure-{number}.png"
        path_.write_bytes(png)
        width, height = _png_size(png)
        display = min(4.0 * inch, 4.0 * inch)
        flowables.append(Spacer(1, 6))
        flowables.append(RLImage(str(path_), width=display, height=display * height / width))
    for choice in sorted(question.choices, key=lambda c: c.label):
        flowables.append(Paragraph(f"<b>{choice.label}.</b>&nbsp;&nbsp;{_pdf_rich(choice.text, workdir, 13.0, _PDF_CONTENT_WIDTH - styles['choice'].leftIndent)}", styles["choice"]))
    if not question.choices:
        response = Table([[""]], colWidths=[_PDF_CONTENT_WIDTH if question.format == "FRQ" else 2.2 * inch], rowHeights=[1.5 * inch if question.format == "FRQ" else 0.32 * inch], hAlign="LEFT")
        response.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, colors.black)]))
        flowables.extend([Spacer(1, 6), response])
    flowables.append(Spacer(1, 10))
    return KeepTogether(flowables)


def build_docx(state: dict, path: Path, answers: bool = False) -> None:
    plan, modules = _ordered(state)
    document = Document()
    title = plan.title + (" — Answer Key" if answers else "")
    document.add_heading(title, level=0)
    document.add_paragraph(f"{saved_profile(state).name} · {assessment_label(plan)} · {plan.total_questions} questions · {plan.module_count} section(s)")
    counter = 0
    for module, slots in modules:
        if not answers:
            document.add_heading(section_name(plan, module), level=2)
            document.add_paragraph(f"{len(slots)} questions · {plan.minutes_per_module[module-1]} minutes")
            section = section_for(plan, module)
            for instruction in section.instructions if section else []:
                _docx_rich(document.add_paragraph(), instruction)
        for slot in slots:
            counter += 1
            if answers:
                paragraph = document.add_paragraph(f"{counter}. ")
                _docx_rich(paragraph, _answer_text(_question(state, slot.id).answer))
                continue
            _docx_question(document, _question(state, slot.id), counter)
        if module != modules[-1][0]:
            document.add_page_break()
    if LOGO_PATH.is_file():
        section = document.add_section(WD_SECTION.NEW_PAGE)
        _center_section_vertically(section)
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(str(LOGO_PATH), width=Inches(LOGO_DISPLAY_INCHES))
    document.save(str(path))


def _center_section_vertically(section) -> None:
    """Set Word's w:vAlign on the section so its lone page centers content top to bottom."""
    sectPr = section._sectPr
    vAlign = sectPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = sectPr.makeelement(qn("w:vAlign"), {})
        # OOXML requires w:vAlign before these trailing sectPr children.
        sectPr.insert_element_before(vAlign, "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")
    vAlign.set(qn("w:val"), "center")


def _docx_rich(paragraph, text: str, size: float = 11.0) -> None:
    for kind, value in _segments(text):
        if kind == "text":
            if value:
                run = paragraph.add_run(value)
                run.font.size = Pt(size)
        else:
            png = math_png(value, fontsize=size)
            width, tall = _png_size(png)
            point_width, point_height = width * 72 / 200, tall * 72 / 200
            if point_width > 6 * 72:
                point_height *= 6 * 72 / point_width
                point_width = 6 * 72
            paragraph.add_run().add_picture(BytesIO(png), width=Pt(point_width), height=Pt(point_height))


def _docx_question(document, question: Question, number: int) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(f"{number}.  ")
    run.bold = True
    _docx_rich(paragraph, question.stem)
    if question.table:
        table = document.add_table(rows=1 + len(question.table.rows), cols=len(question.table.headers))
        table.style = "Table Grid"
        for index, cell in enumerate(question.table.headers):
            _docx_rich(table.cell(0, index).paragraphs[0], cell)
        for row_index, row in enumerate(question.table.rows, start=1):
            for column, cell in enumerate(row):
                _docx_rich(table.cell(row_index, column).paragraphs[0], cell)
    if question.figure:
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(BytesIO(figure_png(question.figure)), width=Inches(3.2))
    for choice in sorted(question.choices, key=lambda c: c.label):
        line = document.add_paragraph()
        line.paragraph_format.left_indent = Inches(0.3)
        run = line.add_run(f"{choice.label}.  ")
        run.bold = True
        _docx_rich(line, choice.text)
    if not question.choices:
        document.add_paragraph("Answer: ____________________")
        if question.format == "FRQ":
            for _ in range(5):
                document.add_paragraph("________________________________________________________________")
    document.add_paragraph()


_LATEX_ESCAPE = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}

# Unicode that the generator may emit as plain text. \ensuremath makes each symbol valid in both text
# and math mode, so one map covers stems, choices and table cells.
_UNICODE = {
    "\u00a0": "~", "\u2009": r"\,", "\u202f": r"\,", "\u200b": "", "\u2010": "-", "\u2011": "-",
    "\u2018": "`", "\u2019": "'", "\u201c": "``", "\u201d": "''", "\u2013": "--", "\u2014": "---", "\u2026": r"\ldots{}",
    "\u00b0": r"\ensuremath{^\circ}", "\u00d7": r"\ensuremath{\times}", "\u00f7": r"\ensuremath{\div}",
    "\u2212": r"\ensuremath{-}", "\u00b1": r"\ensuremath{\pm}", "\u2213": r"\ensuremath{\mp}",
    "\u00b7": r"\ensuremath{\cdot}", "\u22c5": r"\ensuremath{\cdot}",
    "\u2264": r"\ensuremath{\le}", "\u2265": r"\ensuremath{\ge}", "\u2260": r"\ensuremath{\ne}",
    "\u2248": r"\ensuremath{\approx}", "\u2261": r"\ensuremath{\equiv}", "\u221d": r"\ensuremath{\propto}",
    "\u221e": r"\ensuremath{\infty}", "\u221a": r"\ensuremath{\surd}",
    "\u2192": r"\ensuremath{\rightarrow}", "\u2190": r"\ensuremath{\leftarrow}", "\u2194": r"\ensuremath{\leftrightarrow}",
    "\u21d2": r"\ensuremath{\Rightarrow}", "\u21d4": r"\ensuremath{\Leftrightarrow}",
    "\u2208": r"\ensuremath{\in}", "\u2209": r"\ensuremath{\notin}", "\u2205": r"\ensuremath{\emptyset}",
    "\u222a": r"\ensuremath{\cup}", "\u2229": r"\ensuremath{\cap}", "\u2282": r"\ensuremath{\subset}", "\u2286": r"\ensuremath{\subseteq}",
    "\u2200": r"\ensuremath{\forall}", "\u2203": r"\ensuremath{\exists}", "\u2202": r"\ensuremath{\partial}",
    "\u2211": r"\ensuremath{\sum}", "\u220f": r"\ensuremath{\prod}", "\u222b": r"\ensuremath{\int}",
    "\u2220": r"\ensuremath{\angle}", "\u22a5": r"\ensuremath{\perp}", "\u2225": r"\ensuremath{\parallel}",
    "\u2245": r"\ensuremath{\cong}", "\u223c": r"\ensuremath{\sim}", "\u2218": r"\ensuremath{\circ}",
    "\u2227": r"\ensuremath{\wedge}", "\u2228": r"\ensuremath{\vee}", "\u00ac": r"\ensuremath{\neg}",
    "\u2308": r"\ensuremath{\lceil}", "\u2309": r"\ensuremath{\rceil}", "\u230a": r"\ensuremath{\lfloor}", "\u230b": r"\ensuremath{\rfloor}",
    "\u00b2": r"\ensuremath{^2}", "\u00b3": r"\ensuremath{^3}", "\u00b9": r"\ensuremath{^1}",
    "\u2070": r"\ensuremath{^0}", "\u2074": r"\ensuremath{^4}", "\u2075": r"\ensuremath{^5}", "\u2076": r"\ensuremath{^6}",
    "\u2077": r"\ensuremath{^7}", "\u2078": r"\ensuremath{^8}", "\u2079": r"\ensuremath{^9}",
    "\u207a": r"\ensuremath{^+}", "\u207b": r"\ensuremath{^-}", "\u207f": r"\ensuremath{^n}", "\u2071": r"\ensuremath{^i}",
    "\u2080": r"\ensuremath{_0}", "\u2081": r"\ensuremath{_1}", "\u2082": r"\ensuremath{_2}", "\u2083": r"\ensuremath{_3}",
    "\u2084": r"\ensuremath{_4}", "\u2085": r"\ensuremath{_5}", "\u2086": r"\ensuremath{_6}", "\u2087": r"\ensuremath{_7}",
    "\u2088": r"\ensuremath{_8}", "\u2089": r"\ensuremath{_9}",
    "\u00bd": r"\ensuremath{1/2}", "\u00bc": r"\ensuremath{1/4}", "\u00be": r"\ensuremath{3/4}",
    "\u2153": r"\ensuremath{1/3}", "\u2154": r"\ensuremath{2/3}",
}
_GREEK = {"alpha": "\u03b1", "beta": "\u03b2", "gamma": "\u03b3", "delta": "\u03b4", "epsilon": "\u03b5", "zeta": "\u03b6", "eta": "\u03b7", "theta": "\u03b8", "iota": "\u03b9", "kappa": "\u03ba", "lambda": "\u03bb", "mu": "\u03bc", "nu": "\u03bd", "xi": "\u03be", "pi": "\u03c0", "rho": "\u03c1", "sigma": "\u03c3", "tau": "\u03c4", "upsilon": "\u03c5", "phi": "\u03c6", "chi": "\u03c7", "psi": "\u03c8", "omega": "\u03c9", "Delta": "\u0394", "Pi": "\u03a0", "Sigma": "\u03a3", "Omega": "\u03a9"}
for _greek_name, _greek_char in _GREEK.items():
    _UNICODE[_greek_char] = "\\ensuremath{\\" + _greek_name + "}"


def _escape_latex(text: str) -> str:
    output = []
    for character in re.sub(r"\s+", " ", text):
        if character in _UNICODE:
            output.append(_UNICODE[character])
        elif character in _LATEX_ESCAPE:
            output.append(_LATEX_ESCAPE[character])
        else:
            output.append(character)
    return "".join(output)


def _latex_text(text: str) -> str:
    text_html(text)
    def prose(value):
        # Passage/poem line breaks and the boundary before its question carry meaning.
        return r"\newline{}".join(_escape_latex(line) for line in value.split("\n"))
    parts, end = [], 0
    for match in MATH.finditer(text):
        parts.append(prose(text[end:match.start()]))
        # Already-authored LaTeX is copied byte for byte, delimiters included.
        parts.append(match.group(0))
        end = match.end()
    parts.append(prose(text[end:]))
    return "".join(parts)


def _answer_text(answer: str) -> str:
    """Use the bounded expression parser to typeset exact numeric keys in every exporter."""
    import sympy as sp

    text = answer.strip()
    if text in set("ABCDEFGH"):
        return text
    try:
        return r"\(" + sp.latex(ExpressionReader([]).read(text)) + r"\)"
    except (MathUnsupported, TypeError, ValueError):
        # Older saved runs may have text answers. Preserve them rather than guessing notation.
        return text


def _answer_latex(answer: str) -> str:
    return _latex_text(_answer_text(answer))


def _anchor(dx: float, dy: float) -> str:
    parts = []
    if abs(dy) >= 0.35 * (abs(dx) + abs(dy) + 1e-9):
        parts.append("south" if dy > 0 else "north")
    if abs(dx) >= 0.35 * (abs(dx) + abs(dy) + 1e-9):
        parts.append("west" if dx > 0 else "east")
    return " ".join(parts) or "center"


def _tikz_figure(figure: Figure) -> str:
    points = {p.label: p for p in figure.points}
    xs = [p.x for p in figure.points] + [v for c in figure.circles for v in (c.x - c.radius, c.x + c.radius)]
    ys = [p.y for p in figure.points] + [v for c in figure.circles for v in (c.y - c.radius, c.y + c.radius)]
    if figure.show_axes:
        xs.append(0)
        ys.append(0)
    xmin, xmax, ymin, ymax = min(xs), max(xs), min(ys), max(ys)
    scale = min(11.0 / max(xmax - xmin, 1e-6), 7.0 / max(ymax - ymin, 1e-6))
    xmid, ymid = (xmin + xmax) / 2, (ymin + ymax) / 2

    def x(value):
        return (value - xmid) * scale

    def y(value):
        return (value - ymid) * scale

    lines = [r"\begin{center}", r"\begin{tikzpicture}[line width=0.7pt]"]
    if figure.show_axes:
        lines.append(rf"\draw[gray] ({x(xmin):.3f},{y(0):.3f}) -- ({x(xmax):.3f},{y(0):.3f});")
        lines.append(rf"\draw[gray] ({x(0):.3f},{y(ymin):.3f}) -- ({x(0):.3f},{y(ymax):.3f});")
    for segment in figure.segments:
        start, end = points[segment.start], points[segment.end]
        lines.append(rf"\draw ({x(start.x):.3f},{y(start.y):.3f}) -- ({x(end.x):.3f},{y(end.y):.3f});")
    for circle in figure.circles:
        lines.append(rf"\draw ({x(circle.x):.3f},{y(circle.y):.3f}) circle ({circle.radius * scale:.3f});")
    centroid_x = sum(x(p.x) for p in figure.points) / len(figure.points) if figure.points else 0.0
    centroid_y = sum(y(p.y) for p in figure.points) / len(figure.points) if figure.points else 0.0
    for point in figure.points:
        px, py = x(point.x), y(point.y)
        lines.append(rf"\fill ({px:.3f},{py:.3f}) circle (1.3pt) node[anchor={_anchor(px - centroid_x, py - centroid_y)}, inner sep=2.5pt] {{{_escape_latex(point.label)}}};")
    caption = "Coordinates are stated in the question or table." if figure.show_axes else "Figure not drawn to scale."
    lines += [r"\end{tikzpicture}", r"\par\smallskip{\small " + _escape_latex(caption) + r"}", r"\end{center}"]
    return "\n".join(lines)


def _latex_question(question: Question, number: int) -> str:
    lines = [r"\noindent\textbf{" + str(number) + r".}\quad " + _latex_text(question.stem) + r"\par"]
    if question.table:
        columns = len(question.table.headers)
        cell_width = rf"\dimexpr\linewidth/{columns}-2\tabcolsep-2\arrayrulewidth\relax"
        alignment = "|" + (r">{\centering\arraybackslash}p{" + cell_width + "}|") * columns
        rows = [r"\hline " + " & ".join(_latex_text(h) for h in question.table.headers) + r" \\"]
        rows += [r"\hline " + " & ".join(_latex_text(cell) for cell in row) + r" \\" for row in question.table.rows]
        lines.append(r"\begin{center}\begin{tabular}{" + alignment + "}\n" + "\n".join(rows) + "\n" + r"\hline\end{tabular}\end{center}\par")
    if question.figure:
        lines.append(_tikz_figure(question.figure) + r"\par")
    if question.choices:
        lines.append(r"\begin{enumerate}[label=\Alph*., leftmargin=*, itemsep=2pt, topsep=4pt]")
        lines += [r"\item " + _latex_text(choice.text) for choice in sorted(question.choices, key=lambda c: c.label)]
        lines.append(r"\end{enumerate}\par")
    else:
        lines.append(r"\vspace{4pt}\noindent\framebox[\linewidth]{\rule{0pt}{18ex}}\par" if question.format == "FRQ" else r"\vspace{4pt}\noindent\framebox[2.2in]{\rule{0pt}{4ex}}\par")
    lines.append(r"\vspace{14pt}")
    return "\\questionblock{\n" + "\n".join(lines) + "\n}"


def render_latex(state: dict, answers: bool = False) -> str:
    plan, modules = _ordered(state)
    title = plan.title + (" --- Answer Key" if answers else "")
    document = [
        r"\documentclass[11pt]{article}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage{textcomp}",
        r"\usepackage[margin=0.9in]{geometry}",
        r"\usepackage{amsmath,amssymb}",
        r"\usepackage{array}",
        r"\usepackage{tikz}",
        r"\usepackage{graphicx}",
        r"\usepackage{enumitem}",
        r"\setlength{\parindent}{0pt}",
        r"\setlength{\parskip}{6pt}",
        r"\newsavebox{\questionbox}",
        # Fit ordinary questions on one page. Oversize questions remain breakable.
        r"\newcommand{\questionblock}[1]{%",
        r"\sbox{\questionbox}{\begin{minipage}{\linewidth}\setlength{\parskip}{6pt}#1\end{minipage}}%",
        r"\ifdim\dimexpr\ht\questionbox+\dp\questionbox\relax>\textheight #1%",
        r"\else\par\noindent\usebox{\questionbox}\par\fi}",
        r"\begin{document}",
        r"\begin{center}",
        r"{\LARGE\bfseries " + _escape_latex(title) + r"}\\[4pt]",
        r"{\small " + _escape_latex(f"{saved_profile(state).name} - {assessment_label(plan)} - {plan.total_questions} questions - {plan.module_count} section(s)") + r"}",
        r"\end{center}\vspace{6pt}",
    ]
    counter = 0
    for index, (module, slots) in enumerate(modules):
        if not answers:
            document += [r"\section*{" + _escape_latex(section_name(plan, module)) + r"}", r"\hrule\vspace{4pt}", r"{\small " + _escape_latex(f"{len(slots)} questions - {plan.minutes_per_module[module-1]} minutes") + r"}\par", r"\vspace{6pt}"]
            section = section_for(plan, module)
            document += [_latex_text(instruction) + r"\par" for instruction in (section.instructions if section else [])]
        for slot in slots:
            counter += 1
            question = _question(state, slot.id)
            if answers:
                document.append(r"\noindent\textbf{" + str(counter) + r".}\quad " + _answer_latex(question.answer) + r"\par\vspace{6pt}")
            else:
                document.append(_latex_question(question, counter))
        if index != len(modules) - 1:
            document.append(r"\newpage")
    if LOGO_PATH.is_file():
        document += [
            r"\newpage",
            r"\thispagestyle{empty}",
            r"\vspace*{\fill}",
            r"\begin{center}",
            r"\includegraphics[width=0.55\textwidth]{" + LOGO_NAME + r"}",
            r"\end{center}",
            r"\vspace*{\fill}",
        ]
    document.append(r"\end{document}")
    return "\n".join(document) + "\n"


def find_latex() -> str | None:
    for name in ("pdflatex", "xelatex", "lualatex", "tectonic"):
        found = shutil.which(name)
        if found:
            return found
    local = os.environ.get("LOCALAPPDATA")
    candidates = [
        Path(local) / "Programs/MiKTeX/miktex/bin/x64/pdflatex.exe" if local else None,
        Path(r"C:\Program Files\MiKTeX\miktex\bin\x64\pdflatex.exe"),
        Path(r"C:\texlive\2025\bin\windows\pdflatex.exe"),
    ]
    for candidate in candidates:
        try:
            if candidate and candidate.is_file():
                return str(candidate)
        except OSError:
            continue
    return None


def build_latex_pdf(state: dict, path: Path, answers: bool = False) -> None:
    compiler = find_latex()
    if not compiler:
        raise LatexUnavailable("No LaTeX engine found; install MiKTeX or TeX Live")
    path = path.resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    tex = path.parent / f"{path.stem}.tex"
    tex.write_text(render_latex(state, answers), encoding="utf-8")
    if LOGO_PATH.is_file():
        shutil.copyfile(LOGO_PATH, path.parent / LOGO_NAME)
    if Path(compiler).stem.lower() == "tectonic":
        command = [compiler, "--outdir", str(path.parent), str(tex)]
    else:
        command = [compiler, "-no-shell-escape", "-interaction=nonstopmode", "-halt-on-error", "-file-line-error", f"-output-directory={path.parent}", str(tex)]
    result = subprocess.run(command, cwd=path.parent, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=300)
    produced = path.parent / f"{path.stem}.pdf"
    if result.returncode != 0 or not produced.is_file():
        log = path.parent / f"{path.stem}.log"
        detail = log.read_text(encoding="utf-8", errors="replace")[-2000:] if log.is_file() else (result.stdout or "")[-2000:]
        raise RuntimeError(f"LaTeX did not compile {path.name}:\n{detail}")
    for suffix in (".aux", ".log", ".out"):
        (path.parent / f"{path.stem}{suffix}").unlink(missing_ok=True)


def export(state: dict, run_dir: Path, fmt: str = "pdf") -> list[Path]:
    if state.get("status") != "complete":
        raise ValueError("Only a completed test may be exported")
    _ordered(state)  # Fail before creating or overwriting any output file.
    files = RunFiles(run_dir)
    files.prepare(migrate=False)
    written: list[Path] = []
    if fmt in {"pdf", "both"}:
        questions, key = run_dir / "questions.pdf", run_dir / "answer-key.pdf"
        files.sources.mkdir(parents=True, exist_ok=True)
        (files.sources / "questions.tex").write_text(render_latex(state), encoding="utf-8")
        (files.sources / "answer-key.tex").write_text(render_latex(state, answers=True), encoding="utf-8")
        if find_latex():
            try:
                for target, answers in ((questions, False), (key, True)):
                    generated = files.sources / target.name
                    build_latex_pdf(state, generated, answers=answers)
                    generated.replace(target)
            except (RuntimeError, OSError, subprocess.TimeoutExpired) as exc:
                import sys
                print(f"LaTeX export failed, using the built-in renderer instead: {exc}", file=sys.stderr)
                build_pdf(state, questions, image_dir=files.rendering / "questions-images")
                build_pdf(state, key, answers=True, image_dir=files.rendering / "answer-key-images")
        else:
            build_pdf(state, questions, image_dir=files.rendering / "questions-images")
            build_pdf(state, key, answers=True, image_dir=files.rendering / "answer-key-images")
        written += [questions, key]
    if fmt in {"docx", "both"}:
        questions, key = run_dir / "questions.docx", run_dir / "answer-key.docx"
        build_docx(state, questions)
        build_docx(state, key, answers=True)
        written += [questions, key]
    return written
